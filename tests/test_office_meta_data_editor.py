import os
import tempfile
from unittest import TestCase
from unittest.mock import MagicMock, patch
import docx
from PyQt5.QtWidgets import QApplication
from office_meta_data_editor import OfficeMetaDataEditorGUI

from core.error_handler import error_handler


class TestOfficeMetaDataEditor(TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication([])

    def setUp(self):
        self.editor = OfficeMetaDataEditorGUI()
        self.test_dir = tempfile.mkdtemp()
        
    def tearDown(self):
        if hasattr(self, 'test_file') and os.path.exists(self.test_file):
            os.remove(self.test_file)
        os.rmdir(self.test_dir)

    @classmethod
    def tearDownClass(cls):
        cls.app.quit()

    def create_test_document(self, filename="test.docx"):
        """Helper to create a test document with metadata"""
        test_file = os.path.join(self.test_dir, filename)
        doc = docx.Document()
        doc.core_properties.title = "Test Title"
        doc.core_properties.author = "Test Author"
        doc.save(test_file)
        return test_file

    @patch('PyQt5.QtWidgets.QFileDialog.getOpenFileName')
    def test_select_file(self, mock_dialog):
        # Create a test document and mock file dialog
        self.test_file = self.create_test_document()
        mock_dialog.return_value = (self.test_file, '')

        # Test file selection
        self.editor.select()

        # Verify file was loaded
        self.assertEqual(self.editor.current_file, self.test_file)
        self.assertEqual(self.editor.model.rowCount(), 1)
        
        # Verify metadata was loaded
        self.assertEqual(self.editor.title_model.item(0).text(), "Test Title")
        self.assertEqual(self.editor.author_model.item(0).text(), "Test Author")

    def test_set_metadata(self):
        # Create a test document
        self.test_file = self.create_test_document()
        self.editor.current_file = self.test_file

        # Set up models with new metadata
        self.editor.title_model.clear()
        self.editor.author_model.clear()
        self.editor.title_model.appendRow(QStandardItem("New Title"))
        self.editor.author_model.appendRow(QStandardItem("New Author"))

        # Update metadata
        self.editor.set_meta_data()

        # Verify changes were saved
        doc = docx.Document(self.test_file)
        self.assertEqual(doc.core_properties.title, "New Title")
        self.assertEqual(doc.core_properties.author, "New Author")

    def test_set_metadata_no_file(self):
        # Test attempting to set metadata without selecting a file
        self.editor.current_file = None
        with patch('PyQt5.QtWidgets.QMessageBox.warning') as mock_warning:
            self.editor.set_meta_data()
            mock_warning.assert_called_once()

    @patch('PyQt5.QtWidgets.QFileDialog.getOpenFileName')
    def test_select_invalid_file(self, mock_dialog):
        # Test selecting a non-existent file
        mock_dialog.return_value = ("/nonexistent/file.docx", '')
        with patch('PyQt5.QtWidgets.QMessageBox.critical') as mock_critical:
            self.editor.select()
            mock_critical.assert_called_once()

    def test_empty_metadata(self):
        # Create document with empty metadata
        self.test_file = os.path.join(self.test_dir, "empty.docx")
        doc = docx.Document()
        doc.save(self.test_file)
        
        self.editor.current_file = self.test_file
        
        # Mock file dialog
        with patch('PyQt5.QtWidgets.QFileDialog.getOpenFileName', return_value=(self.test_file, '')):
            self.editor.select()
            
            # Verify empty metadata is handled correctly
            self.assertEqual(self.editor.title_model.item(0).text(), "")
            self.assertEqual(self.editor.author_model.item(0).text(), "")

    def test_initial_ui_state(self):
        """Test initial state of UI components"""
        # Test initial window properties
        self.assertTrue(self.editor.isVisible())
        
        # Test ListView models are initialized empty
        self.assertEqual(self.editor.model.rowCount(), 0)
        self.assertEqual(self.editor.title_model.rowCount(), 0)
        self.assertEqual(self.editor.author_model.rowCount(), 0)
        self.assertEqual(self.editor.created_model.rowCount(), 0)
        self.assertEqual(self.editor.modified_model.rowCount(), 0)
        
        # Verify ListView widgets have correct models
        self.assertEqual(self.editor.select_ListView.model(), self.editor.model)
        self.assertEqual(self.editor.title_ListView.model(), self.editor.title_model)
        self.assertEqual(self.editor.author_ListView.model(), self.editor.author_model)
        self.assertEqual(self.editor.created_ListView.model(), self.editor.created_model)
        self.assertEqual(self.editor.modified_ListView.model(), self.editor.modified_model)

    def test_menu_actions(self):
        """Test menu actions are properly connected"""
        # Mock the select and close methods
        self.editor.select = MagicMock()
        self.editor.close = MagicMock()
        
        # Trigger menu actions
        self.editor.actionSelect_Files.trigger()
        self.editor.actionExit.trigger()
        
        # Verify methods were called
        self.editor.select.assert_called_once()
        self.editor.close.assert_called_once()

    def test_apply_changes_button_interaction(self):
        """Test apply changes button interaction"""
        # Create test document and set it up
        self.test_file = self.create_test_document()
        self.editor.current_file = self.test_file
        
        # Mock the set_meta_data method
        self.editor.set_meta_data = MagicMock()
        
        # Click the apply changes button
        self.editor.apply_changes_button.click()
        
        # Verify the method was called
        self.editor.set_meta_data.assert_called_once()

    def test_listview_interaction(self):
        """Test ListView interactions and updates"""
        # Create and set up test document
        self.test_file = self.create_test_document()
        self.editor.current_file = self.test_file
        
        # Simulate file selection
        with patch('PyQt5.QtWidgets.QFileDialog.getOpenFileName', 
                  return_value=(self.test_file, '')):
            self.editor.select()
        
        # Test that all ListViews are populated
        self.assertEqual(self.editor.model.rowCount(), 1)  # File list
        self.assertEqual(self.editor.title_model.rowCount(), 1)
        self.assertEqual(self.editor.author_model.rowCount(), 1)
        self.assertEqual(self.editor.created_model.rowCount(), 1)
        self.assertEqual(self.editor.modified_model.rowCount(), 1)
        
        # Test clearing models
        self.editor.model.clear()
        self.assertEqual(self.editor.model.rowCount(), 0)
        
        # Test model data updates
        self.editor.title_model.clear()
        self.editor.title_model.appendRow(QStandardItem("Updated Title"))
        self.assertEqual(self.editor.title_model.item(0).text(), "Updated Title")
